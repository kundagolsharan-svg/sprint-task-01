import sqlite3
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]


def test_screener_has_six_presets():
    assert len(openpyxl.load_workbook(ROOT / "output" / "screener_output.xlsx", read_only=True).sheetnames) == 6


def test_peer_comparison_has_eleven_groups():
    assert len(openpyxl.load_workbook(ROOT / "output" / "peer_comparison.xlsx", read_only=True).sheetnames) == 11


def test_peer_percentiles_have_all_rows():
    with sqlite3.connect(ROOT / "nifty100.db") as db:
        assert db.execute("SELECT COUNT(*) FROM peer_percentiles").fetchone()[0] == 560


def test_foreign_keys_are_clean():
    with sqlite3.connect(ROOT / "nifty100.db") as db:
        assert db.execute("PRAGMA foreign_key_check").fetchall() == []


def test_radar_chart_count():
    assert len(list((ROOT / "reports" / "radar_charts").glob("*.png"))) == 56
