"""Peer-group percentile rankings, Excel comparison, and radar charts."""

import sqlite3
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[2]
PEER_SOURCE = ROOT / "tests" / "etl" / "1788501620796-5060f580-peer_groups.xlsx"
METRICS = {
    "ROE": "return_on_equity_pct", "ROCE": "return_on_capital_employed_pct", "NPM": "net_profit_margin_pct",
    "DE": "debt_to_equity", "FCF": "free_cash_flow_cr", "PAT CAGR 5yr": "pat_cagr_5yr",
    "Revenue CAGR 5yr": "revenue_cagr_5yr", "EPS CAGR 5yr": "eps_cagr_5yr", "Interest Coverage": "interest_coverage", "Asset Turnover": "asset_turnover",
}


def latest_ratios(db):
    frame = pd.read_sql_query("SELECT * FROM financial_ratios", db)
    frame = frame[frame.year != "TTM"].sort_values("year").groupby("company_id", as_index=False).tail(1)
    return frame


def populate_peer_percentiles(db_path=ROOT / "nifty100.db"):
    membership = pd.read_excel(PEER_SOURCE)
    db = sqlite3.connect(db_path)
    ratios = latest_ratios(db)
    merged = membership.merge(ratios, on="company_id", how="left")
    rows = []
    for (group, year), group_frame in merged.groupby(["peer_group_name", "year"], dropna=False):
        for label, column in METRICS.items():
            values = pd.to_numeric(group_frame[column], errors="coerce")
            ranks = values.rank(pct=True, method="average")
            if label == "DE":
                ranks = 1 - ranks
            for index, item in group_frame.iterrows():
                rows.append((item.company_id, group, label, item[column], None if pd.isna(ranks[index]) else float(ranks[index]), str(year)))
    db.execute("DROP TABLE IF EXISTS peer_percentiles")
    db.execute("CREATE TABLE peer_percentiles(company_id TEXT NOT NULL, peer_group_name TEXT NOT NULL, metric TEXT NOT NULL, value REAL, percentile_rank REAL, year TEXT NOT NULL, PRIMARY KEY(company_id, peer_group_name, metric, year), FOREIGN KEY(company_id) REFERENCES companies(company_id))")
    db.executemany("INSERT INTO peer_percentiles VALUES (?, ?, ?, ?, ?, ?)", rows)
    db.commit()
    comparison = merged.copy()
    comparison = comparison.merge(pd.DataFrame(rows, columns=["company_id", "peer_group_name", "metric", "value", "percentile_rank", "rank_year"]).pivot_table(index=["company_id", "peer_group_name"], columns="metric", values="percentile_rank", aggfunc="first").reset_index(), on=["company_id", "peer_group_name"], how="left", suffixes=("", "_percentile"))
    output = ROOT / "output" / "peer_comparison.xlsx"
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for group, group_frame in comparison.groupby("peer_group_name"):
            group_frame.to_excel(writer, sheet_name=str(group)[:31], index=False)
    workbook = __import__("openpyxl").load_workbook(output)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, (int, float)) and 0 <= cell.value <= 1:
                    cell.fill = PatternFill("solid", fgColor="C6EFCE" if cell.value >= .75 else "FFEB9C" if cell.value >= .25 else "FFC7CE")
    workbook.save(output)
    db.close()
    return len(rows), comparison


def generate_radar_charts(db_path=ROOT / "nifty100.db"):
    membership = pd.read_excel(PEER_SOURCE)
    db = sqlite3.connect(db_path)
    ratios = latest_ratios(db)
    data = membership.merge(ratios, on="company_id", how="left")
    target = ROOT / "reports" / "radar_charts"
    target.mkdir(parents=True, exist_ok=True)
    axes = ["return_on_equity_pct", "return_on_capital_employed_pct", "net_profit_margin_pct", "debt_to_equity", "free_cash_flow_cr", "pat_cagr_5yr", "revenue_cagr_5yr", "composite_quality_score"]
    labels = ["ROE", "ROCE", "NPM", "D/E", "FCF", "PAT CAGR", "Revenue CAGR", "Composite"]
    for _, item in data.iterrows():
        group_frame = data[data.peer_group_name == item.peer_group_name]
        values = [float(pd.to_numeric(item[axis], errors="coerce") or 0) for axis in axes]
        average = [float(pd.to_numeric(group_frame[axis], errors="coerce").mean() or 0) for axis in axes]
        angles = np.linspace(0, 2 * np.pi, len(labels), endpoint=False).tolist()
        values += values[:1]; average += average[:1]; angles += angles[:1]
        figure, polar = plt.subplots(figsize=(7, 7), subplot_kw={"polar": True})
        polar.plot(angles, values, linewidth=2); polar.fill(angles, values, alpha=.2)
        polar.plot(angles, average, linestyle="--", linewidth=1.5)
        polar.set_xticks(angles[:-1]); polar.set_xticklabels(labels, fontsize=8); polar.set_title(f"{item.company_id} - {item.peer_group_name}")
        figure.savefig(target / f"{item.company_id}_radar.png", dpi=140, bbox_inches="tight"); plt.close(figure)
    db.close()


if __name__ == "__main__":
    print("peer percentile rows:", populate_peer_percentiles()[0]); generate_radar_charts()
