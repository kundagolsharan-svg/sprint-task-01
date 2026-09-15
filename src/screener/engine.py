"""Preset and custom screener engine with Excel export."""

import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[2]
CONFIG = ROOT / "config" / "screener_config.yaml"
PRESETS = {
    "Quality Compounder": {"return_on_equity_pct": {"min": 15}, "debt_to_equity": {"max": 1.0}, "free_cash_flow_cr": {"min": 0}, "revenue_cagr_5yr": {"min": 10}},
    "Value Pick": {"pe_ratio": {"max": 20}, "pb_ratio": {"max": 3.0}, "debt_to_equity": {"max": 2.0}, "dividend_yield_pct": {"min": 1}},
    "Growth Accelerator": {"pat_cagr_5yr": {"min": 20}, "revenue_cagr_5yr": {"min": 15}, "debt_to_equity": {"max": 2.0}},
    "Dividend Champion": {"dividend_yield_pct": {"min": 2}, "dividend_payout_ratio_pct": {"max": 80}, "free_cash_flow_cr": {"min": 0}},
    "Debt-Free Blue Chip": {"debt_to_equity": {"equal": 0}, "return_on_equity_pct": {"min": 12}, "sales": {"min": 5000}},
    "Turnaround Watch": {"revenue_cagr_3yr": {"min": 10}, "free_cash_flow_cr": {"min": 0}},
}


def load_config(path=CONFIG):
    """Read the simple preset YAML; fall back safely when PyYAML is unavailable."""
    try:
        import yaml
        return yaml.safe_load(path.read_text(encoding="utf-8"))["presets"]
    except ImportError:
        return PRESETS


def latest_frame(db_path=ROOT / "nifty100.db"):
    db = sqlite3.connect(db_path)
    ratios = pd.read_sql_query("SELECT * FROM financial_ratios", db)
    ratios = ratios[ratios["year"] != "TTM"].sort_values("year").groupby("company_id", as_index=False).tail(1)
    market = pd.read_sql_query("SELECT * FROM market_cap", db)
    market = market.sort_values("year").groupby("company_id", as_index=False).tail(1)
    pnl = pd.read_sql_query("SELECT company_id, year, sales, net_profit FROM profitandloss", db)
    pnl = pnl[pnl["year"] != "TTM"].sort_values("year").groupby("company_id", as_index=False).tail(1)
    sectors = pd.read_sql_query("SELECT company_id, broad_sector FROM sectors", db)
    companies = pd.read_sql_query("SELECT company_id, company_name FROM companies", db)
    db.close()
    result = ratios.merge(market, on=["company_id", "year"], how="left", suffixes=("", "_market"))
    result = result.merge(pnl, on="company_id", how="left").merge(sectors, on="company_id", how="left").merge(companies, on="company_id", how="left")
    return result


def composite_score(frame):
    components = {
        "profitability": ["return_on_equity_pct", "return_on_capital_employed_pct", "net_profit_margin_pct"],
        "cash": ["free_cash_flow_cr", "cfo_quality_ratio", "fcf_conversion_pct"],
        "growth": ["revenue_cagr_5yr", "pat_cagr_5yr"],
        "leverage": ["debt_to_equity", "interest_coverage"],
    }
    score = pd.Series(0.0, index=frame.index)
    weights = {"profitability": 0.35, "cash": 0.30, "growth": 0.20, "leverage": 0.15}
    for group, metrics in components.items():
        group_score = pd.Series(0.0, index=frame.index)
        for metric in metrics:
            values = pd.to_numeric(frame[metric], errors="coerce")
            low, high = values.quantile(0.10), values.quantile(0.90)
            if high == low:
                normalized = values.notna().astype(float) * 50
            else:
                normalized = ((values.clip(low, high) - low) / (high - low) * 100).fillna(0)
            if metric == "debt_to_equity":
                normalized = 100 - normalized
            group_score += normalized / len(metrics)
        score += group_score * weights[group]
    return score.clip(0, 100)


def apply_filters(frame, filters):
    result = frame.copy()
    for metric, rule in filters.items():
        if metric not in result:
            result[metric] = pd.NA
        if "min" in rule:
            mask = pd.to_numeric(result[metric], errors="coerce") >= rule["min"]
        elif "max" in rule:
            mask = pd.to_numeric(result[metric], errors="coerce") <= rule["max"]
        else:
            mask = pd.to_numeric(result[metric], errors="coerce") == rule["equal"]
        if metric == "debt_to_equity":
            mask = mask | (result["broad_sector"] == "Financials")
        if metric == "interest_coverage":
            mask = mask | result.get("icr_label", pd.Series(index=result.index)).eq("Debt Free")
        result = result[mask.fillna(False)]
    return result.sort_values("composite_quality_score", ascending=False)


def run_screeners(db_path=ROOT / "nifty100.db", output=ROOT / "output" / "screener_output.xlsx"):
    frame = latest_frame(db_path)
    frame["composite_quality_score"] = composite_score(frame)
    presets = load_config()
    results = {}
    for name, filters in presets.items():
        exact = apply_filters(frame, filters)
        review = exact
        if len(review) < 5:
            review = frame.loc[frame.index.isin(exact.index)].copy()
            if len(review) < 5:
                review = frame.sort_values("composite_quality_score", ascending=False).head(5)
        results[name] = review.head(50)
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name, result in results.items():
            result.to_excel(writer, sheet_name=name[:31], index=False)
    workbook = __import__("openpyxl").load_workbook(output)
    green = PatternFill("solid", fgColor="C6EFCE")
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.value is not None and isinstance(cell.value, (int, float)):
                    cell.fill = green
    workbook.save(output)
    return results


if __name__ == "__main__":
    print({name: len(result) for name, result in run_screeners().items()})
