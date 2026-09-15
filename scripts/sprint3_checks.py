"""Verify Sprint 3 screener and peer-engine exit criteria."""

import sqlite3
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
PRESETS = ["Quality Compounder", "Value Pick", "Growth Accelerator", "Dividend Champion", "Debt-Free Blue Chip", "Turnaround Watch"]


def main():
    screener = openpyxl.load_workbook(ROOT / "output" / "screener_output.xlsx", read_only=True)
    peer = openpyxl.load_workbook(ROOT / "output" / "peer_comparison.xlsx", read_only=True)
    with sqlite3.connect(ROOT / "nifty100.db") as db:
        percentile_rows = db.execute("SELECT COUNT(*) FROM peer_percentiles").fetchone()[0]
        groups = db.execute("SELECT COUNT(DISTINCT peer_group_name) FROM peer_percentiles").fetchone()[0]
        foreign_keys = len(db.execute("PRAGMA foreign_key_check").fetchall())
        quality = db.execute("SELECT COUNT(*) FROM financial_ratios r WHERE r.year = (SELECT MAX(r2.year) FROM financial_ratios r2 WHERE r2.company_id = r.company_id AND r2.year <> 'TTM') AND r.return_on_equity_pct > 15 AND r.debt_to_equity < 1").fetchone()[0]
    preset_sizes = {sheet: screener[sheet].max_row - 1 for sheet in screener.sheetnames}
    checks = {
        "six_preset_sheets": screener.sheetnames == PRESETS,
        "preset_sizes_5_to_50": all(5 <= size <= 50 for size in preset_sizes.values()),
        "eleven_peer_sheets": len(peer.sheetnames) == 11,
        "peer_percentile_rows": percentile_rows == 560,
        "eleven_peer_groups": groups == 11,
        "radar_charts": len(list((ROOT / "reports" / "radar_charts").glob("*.png"))) == 56,
        "quality_filter": 5 <= quality <= 50,
        "foreign_key_check": foreign_keys == 0,
    }
    for name, passed in checks.items():
        print(f"{name}: {'PASS' if passed else 'FAIL'}")
    print("preset_sizes=", preset_sizes)
    print("percentile_rows=", percentile_rows)
    return 0 if all(checks.values()) else 1


if __name__ == "__main__":
    raise SystemExit(main())
