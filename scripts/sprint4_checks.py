"""Verify Sprint 4 dashboard and valuation exit criteria."""

import sqlite3
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
PAGES = [f"{index:02d}_{name}.py" for index, name in enumerate(["home", "profile", "screener", "peers", "trends", "sectors", "capital", "reports"], 1)]
VALUATION_COLUMNS = ["company_id", "company_name", "sector", "P/E", "P/B", "EV/EBITDA", "fcf_yield_pct", "5yr_median_PE", "PE_vs_sector_median_pct", "flag"]


def main():
    summary = openpyxl.load_workbook(ROOT / "output" / "valuation_summary.xlsx", read_only=True)
    headers = next(summary.active.iter_rows(values_only=True))
    with sqlite3.connect(ROOT / "nifty100.db") as db:
        valuation_rows = db.execute("SELECT COUNT(*) FROM valuation_summary").fetchone()[0]
        foreign_keys = len(db.execute("PRAGMA foreign_key_check").fetchall())
    checks = {
        "eight_page_files": all((ROOT / "src" / "dashboard" / "pages" / page).exists() for page in PAGES),
        "valuation_rows_92": summary.active.max_row - 1 == 92,
        "valuation_columns": list(headers) == VALUATION_COLUMNS,
        "valuation_table_rows": valuation_rows == 92,
        "valuation_flags_csv": (ROOT / "output" / "valuation_flags.csv").exists(),
        "foreign_key_check": foreign_keys == 0,
    }
    for name, passed in checks.items():
        print(f"{name}: {'PASS' if passed else 'FAIL'}")
    return 0 if all(checks.values()) else 1


if __name__ == "__main__":
    raise SystemExit(main())
