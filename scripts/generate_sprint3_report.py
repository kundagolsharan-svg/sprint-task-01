"""Generate the Sprint 3 completion report."""

import sqlite3
from pathlib import Path
import subprocess
import sys

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "sprint3_completion_report.md"


def main():
    tests = subprocess.run([sys.executable, "-m", "pytest", "-q"], cwd=ROOT, capture_output=True, text=True)
    checks = subprocess.run([sys.executable, "scripts/sprint3_checks.py"], cwd=ROOT, capture_output=True, text=True)
    with sqlite3.connect(ROOT / "nifty100.db") as db:
        ratios = db.execute("SELECT COUNT(*) FROM financial_ratios").fetchone()[0]
        percentiles = db.execute("SELECT COUNT(*) FROM peer_percentiles").fetchone()[0]
        groups = db.execute("SELECT COUNT(DISTINCT peer_group_name) FROM peer_percentiles").fetchone()[0]
        foreign_keys = len(db.execute("PRAGMA foreign_key_check").fetchall())
    screener = openpyxl.load_workbook(ROOT / "output" / "screener_output.xlsx", read_only=True)
    peer = openpyxl.load_workbook(ROOT / "output" / "peer_comparison.xlsx", read_only=True)
    lines = [
        "# Sprint 3 Completion Report", "", "## Definition of Done", "",
        f"- Screener workbook sheets: **{len(screener.sheetnames)}** (required: 6)",
        f"- Peer comparison workbook sheets: **{len(peer.sheetnames)}** (required: 11)",
        f"- Peer percentile rows: **{percentiles}**", f"- Peer groups: **{groups}** (required: 11)",
        f"- Financial ratio rows: **{ratios}**", f"- Radar charts: **{len(list((ROOT / 'reports' / 'radar_charts').glob('*.png')))}**",
        f"- Foreign-key errors: **{foreign_keys}**", f"- Test command exit code: **{tests.returncode}**", f"- Sprint 3 check exit code: **{checks.returncode}**", "",
        "## Preset Sizes", "", "```text", checks.stdout.strip(), "```", "",
        "## Test Output", "", "```text", tests.stdout.strip(), "```", "",
        "## Deliverables", "", "- `config/screener_config.yaml`", "- `src/screener/engine.py`", "- `src/analytics/peer.py`", "- `output/screener_output.xlsx`", "- `output/peer_comparison.xlsx`", "- `reports/radar_charts/`", "- `peer_percentiles` SQLite table", "",
    ]
    OUTPUT.write_text("\n".join(lines), encoding="utf-8")
    print(OUTPUT)
    return tests.returncode or checks.returncode


if __name__ == "__main__":
    raise SystemExit(main())
